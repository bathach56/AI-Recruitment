import io
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

import pymupdf
from docx import Document

LAB = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(LAB))
from core.document_parser import MAX_BYTES, ParseError, parse_bytes, parse_file


class ParserTests(unittest.TestCase):
    def assert_code(self, code, data, name):
        with self.assertRaises(ParseError) as raised:
            parse_bytes(data, name)
        self.assertEqual(raised.exception.code, code)

    def test_docx_order_unicode_and_nested_table(self):
        doc = Document()
        doc.add_paragraph('Nguyễn Văn An — C#, C++, .NET')
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = 'SQL Server'
        cell.add_table(rows=1, cols=1).cell(0, 0).text = 'Docker'
        doc.add_paragraph('5.5 năm kinh nghiệm')
        buffer = io.BytesIO()
        doc.save(buffer)
        result = parse_bytes(buffer.getvalue(), '../../cv.docx')
        self.assertEqual(result['source_name'], 'cv.docx')
        self.assertEqual(result['raw_text'].splitlines(),
                         ['Nguyễn Văn An — C#, C++, .NET', 'SQL Server', 'Docker', '5.5 năm kinh nghiệm'])
        self.assertIn('/cell:1', result['segments'][1]['location'])

    def test_all_sample_docx(self):
        paths = list((LAB.parent / 'data/sample_cvs').glob('*.docx'))
        self.assertEqual(len(paths), 20)
        for path in paths:
            with self.subTest(path=path.name):
                result = parse_file(path)
                self.assertGreater(len(result['raw_text']), 100)
                self.assertEqual(result['raw_text'], '\n'.join(s['text'] for s in result['segments']))

    def test_pdf_pages_and_blank_warning(self):
        with pymupdf.open() as doc:
            doc.new_page().insert_text((72, 72), 'C# ASP.NET Core - 5 years')
            doc.new_page()
            doc.new_page().insert_text((72, 72), 'Docker SQL Server')
            result = parse_bytes(doc.tobytes(), 'cv.pdf')
        self.assertEqual([s['location'] for s in result['segments']], ['page:1', 'page:3'])
        self.assertIn('C# ASP.NET', result['raw_text'])
        self.assertTrue(any('Trang 2' in w for w in result['warnings']))

    def test_empty_pdf_and_docx(self):
        with pymupdf.open() as doc:
            doc.new_page()
            self.assert_code('no_text', doc.tobytes(), 'empty.pdf')
        buffer = io.BytesIO()
        Document().save(buffer)
        self.assert_code('no_text', buffer.getvalue(), 'empty.docx')

    def test_encrypted_pdf(self):
        with pymupdf.open() as doc:
            doc.new_page().insert_text((72, 72), 'Private')
            data = doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256,
                              owner_pw='owner', user_pw='test-password')
        self.assert_code('encrypted_pdf', data, 'locked.pdf')

    def test_bad_inputs(self):
        self.assert_code('unsupported_type', b'data', 'cv.doc')
        self.assert_code('empty_file', b'', 'cv.pdf')
        self.assert_code('file_too_large', b'x' * (MAX_BYTES + 1), 'cv.pdf')
        self.assert_code('invalid_file', b'not a PDF', 'cv.pdf')
        self.assert_code('invalid_file', b'broken ZIP', 'cv.docx')

    def test_missing_dependency_no_fallback(self):
        with patch.dict(sys.modules, {'pymupdf': None}):
            self.assert_code('missing_dependency', b'%PDF-1.7\n', 'cv.pdf')

    def test_archive_limit(self):
        doc = Document()
        doc.add_paragraph('example')
        buffer = io.BytesIO()
        doc.save(buffer)
        with patch('core.document_parser.MAX_UNPACKED', 1):
            self.assert_code('archive_too_large', buffer.getvalue(), 'cv.docx')


if __name__ == '__main__':
    unittest.main()
